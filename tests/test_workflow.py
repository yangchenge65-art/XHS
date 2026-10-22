import io
import tempfile
import unittest
from datetime import datetime, timedelta
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from hotdesk.core import Store, Service, SourceError, export_excel
from hotdesk.web import create_app
from hotdesk.scheduler import tick

TZ = ZoneInfo('Asia/Shanghai')


class Source:
    cost = '0.50'
    def __init__(self):
        self.calls = 0
        self.fail = False

    def ready(self, platform, kind):
        return None

    def fetch(self, platform, keyword=None):
        self.calls += 1
        if self.fail:
            raise SourceError('测试数据源不可用')
        return [dict(id='one', title='KPL 决赛名场面', content='KPL 决赛讨论',
                     url='https://www.douyin.com/video/123', published_at='2026-10-21T12:00:00+08:00',
                     rank=2, heat=10000),
                dict(id='old', title='KPL 旧闻', content='', url='',
                     published_at='2026-09-01T00:00:00+08:00', rank=5, heat=1)]


class Workflow(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.now = datetime(2026, 10, 22, 10, tzinfo=TZ)
        self.store = Store(self.tmp.name + '/test.db')
        self.source = Source()
        self.service = Service(self.store, self.source, now=lambda: self.now)

    def tearDown(self):
        self.tmp.cleanup()

    def test_daily_is_idempotent_and_counts_distinct_days(self):
        self.service.collect()
        self.service.collect()
        self.assertEqual(self.source.calls, 2)
        self.now += timedelta(days=1)
        self.service.collect()
        rows = self.service.report()['items']
        self.assertTrue(rows)
        self.assertEqual(rows[0]['days'], 2)
        self.assertIn('2 天', rows[0]['reason'])
        self.assertEqual(self.store.spent(self.now), 2.0)

    def test_keyword_seven_days_and_cache(self):
        first = self.service.search(' KPL ')
        second = self.service.search('kpl')
        self.assertEqual(len(first['items']), 2)
        self.assertTrue(all('旧闻' not in i['title'] for i in first['items']))
        self.assertTrue(second['cached'])
        self.assertEqual(self.source.calls, 2)
        self.assertEqual(self.service.report()['items'], [])

    def test_failed_search_is_not_cached_as_empty_success(self):
        self.source.fail = True
        result = self.service.search('肖战')
        self.assertEqual(result['status'], 'error')
        self.assertEqual(len(result['errors']), 2)
        self.source.fail = False
        self.assertEqual(self.service.search('肖战')['status'], 'ok')

    def test_budget_reservation_is_persistent_and_resets_by_shanghai_month(self):
        self.assertTrue(self.store.reserve(self.now, '29.5'))
        self.service.collect()
        self.assertEqual(self.source.calls, 1)
        self.assertEqual(self.store.spent(self.now), 30)
        self.assertFalse(self.store.reserve(self.now, '0.01'))
        next_month = self.now.replace(month=11)
        self.assertTrue(self.store.reserve(next_month, '0.5'))
        self.assertEqual(self.store.spent(next_month), .5)
        self.assertEqual(Store(self.tmp.name + '/test.db').spent(self.now), 30)

    def test_excel_has_required_columns_and_never_executes_formula(self):
        report = self.service.search('KPL')
        report['items'][0]['title'] = '=HYPERLINK("bad")'
        book = load_workbook(io.BytesIO(export_excel(report, demo=True)))
        sheet = book['热点素材']
        self.assertEqual([c.value for c in sheet[1]][:3], ['热点关键词', '热点具体内容', '热点相关链接'])
        self.assertEqual(sheet['A2'].data_type, 's')
        self.assertEqual(sheet['A2'].value, '=HYPERLINK("bad")')
        self.assertIn('示例', book['采集说明']['B1'].value)

    def test_web_real_empty_demo_isolated_and_export(self):
        app = create_app(self.tmp.name, provider=self.source, now=lambda: self.now)
        client = app.test_client()
        self.assertEqual(client.get('/').status_code, 200)
        self.assertEqual(client.get('/api/report').json['items'], [])
        demo = client.get('/api/report?demo=1').json
        self.assertTrue(demo['demo'])
        self.assertGreater(len(demo['items']), 0)
        self.assertEqual(client.get('/api/report').json['items'], [])
        self.assertEqual(client.post('/api/search', json={'keyword': ''}).status_code, 400)
        out = client.get('/api/export?demo=1')
        self.assertEqual(out.status_code, 200)
        self.assertTrue(load_workbook(io.BytesIO(out.data)).worksheets)
        self.assertEqual(client.post('/api/collect', headers={'Origin': 'https://evil.example'}).status_code, 403)

    def test_scheduler_waits_until_nine_and_does_not_repeat_failed_calls(self):
        self.now = self.now.replace(hour=8)
        tick(self.service)
        self.assertEqual(self.source.calls, 0)
        self.now = self.now.replace(hour=9)
        self.source.fail = True
        tick(self.service)
        tick(self.service)
        self.assertEqual(self.source.calls, 2)
        self.now += timedelta(days=1)
        self.source.fail = False
        tick(self.service)
        self.assertEqual(self.source.calls, 4)

    def test_partial_search_can_export_without_paid_retry(self):
        original = self.source.ready
        self.source.ready = lambda platform, kind: '暂不可用' if platform == 'xiaohongshu' else original(platform, kind)
        app = create_app(self.tmp.name, provider=self.source, now=lambda: self.now)
        client = app.test_client()
        result = client.post('/api/search', json={'keyword': 'KPL'})
        self.assertEqual(result.json['status'], 'partial')
        self.assertEqual(self.source.calls, 1)
        exported = client.get('/api/export?keyword=KPL')
        self.assertEqual(exported.status_code, 200)
        self.assertEqual(self.source.calls, 1)
        workbook = load_workbook(io.BytesIO(exported.data))
        self.assertEqual(workbook['热点素材'].max_row, 2)
        self.assertTrue(any('暂不可用' in str(row[1].value) for row in workbook['采集说明']))

    def test_keyword_includes_observed_historical_hot_list(self):
        self.service.collect()
        self.source.fail = True
        result = self.service.search('KPL')
        self.assertEqual(result['status'], 'partial')
        self.assertTrue(result['items'])
        self.assertTrue(all(i['kind'] == '榜单记录' for i in result['items']))
        self.assertEqual(self.service.search('不存在的词')['items'], [])


if __name__ == '__main__':
    unittest.main()
