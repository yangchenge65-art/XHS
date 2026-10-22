"""Optional browser check: use system Python with Playwright and Chromium installed."""
from pathlib import Path
import os
from uuid import uuid4
from playwright.sync_api import sync_playwright, expect
from openpyxl import load_workbook


def run():
    artifacts = Path(__file__).resolve().parents[1] / 'data'
    artifacts.mkdir(exist_ok=True)
    errors = []
    with sync_playwright() as p:
        browser = p.chromium.launch(executable_path='/usr/bin/chromium', headless=True, args=['--no-sandbox'])
        page = browser.new_page(viewport={'width': 1440, 'height': 1000})
        page.on('pageerror', lambda error: errors.append(str(error)))
        page.goto(os.getenv('HOTDESK_TEST_URL', 'http://127.0.0.1:8000'), wait_until='networkidle')
        expect(page.locator('#empty')).to_be_visible()
        page.locator('#demo-toggle').check()
        expect(page.locator('#results tr')).to_have_count(9)
        page.screenshot(path=str(artifacts / 'desktop-demo.png'), full_page=True)
        page.locator('[data-platform="douyin"]').click()
        expect(page.locator('#results tr')).to_have_count(4)
        page.locator('[data-platform="xiaohongshu"]').click()
        expect(page.locator('#results tr')).to_have_count(3)
        page.locator('[data-platform="news"]').click()
        expect(page.locator('#results tr')).to_have_count(2)
        page.locator('[data-platform="all"]').click()
        before = page.locator('#results tr').all_text_contents()
        page.locator('.replace-button').first.click()
        expect(page.locator('#results tr').first).not_to_have_text(before[0])
        after = page.locator('#results tr').all_text_contents()
        assert after[1:] == before[1:]
        with page.expect_download() as download:
            page.locator('#export-button').click()
        daily_path = artifacts / '热点素材-每日九条-示例.xlsx'
        download.value.save_as(str(daily_path))
        daily_book = load_workbook(daily_path)
        assert daily_book['热点素材'].max_row == 10
        assert '示例' in daily_book['采集说明']['B1'].value
        page.locator('[data-view="search"]').click()
        page.locator('#keyword').fill('浏览器验证' + uuid4().hex[:8])
        page.locator('#search-button').click()
        expect(page.locator('#results tr')).to_have_count(6)
        with page.expect_download() as download:
            page.locator('#export-button').click()
        path = artifacts / '热点素材-示例.xlsx'
        download.value.save_as(str(path))
        book = load_workbook(path)
        assert book['热点素材'].max_row == 7
        assert '示例' in book['采集说明']['B1'].value
        page.locator('[data-view="settings"]').click()
        expect(page.locator('.source-card')).to_have_count(2)
        expect(page.locator('#openai-key')).to_have_attribute('type', 'password')
        expect(page.locator('#engine-select')).to_have_value('codex')
        expect(page.locator('#codex-panel')).to_be_visible()
        page.locator('#engine-select').select_option('openai')
        expect(page.locator('#api-config')).to_be_visible()
        expect(page.locator('#openai-state')).to_contain_text('尚未连接')
        page.locator('#engine-select').select_option('codex')
        expect(page.locator('#codex-panel')).to_be_visible()
        page.locator('#demo-toggle').uncheck()
        page.locator('[data-view="daily"]').click()
        expect(page.locator('#results tr')).to_have_count(0)
        expect(page.locator('#demo-banner')).to_be_hidden()
        page.set_viewport_size({'width': 390, 'height': 844})
        page.screenshot(path=str(artifacts / 'mobile-real.png'), full_page=True)
        assert page.evaluate('() => document.documentElement.scrollWidth <= window.innerWidth')
        browser.close()
    assert not errors, errors
    print('PASS: 4/3/2 quotas, single replacement, demo/real isolation, filtering, search, final Excel, account UI, mobile, no JS errors')


if __name__ == '__main__':
    run()
