import io
import tempfile
import unittest
from types import SimpleNamespace
from collections import Counter
from datetime import datetime, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from hotdesk.core import Store, SourceError
from hotdesk.briefs import Briefing
from hotdesk.web import create_app
from hotdesk.scheduler import tick
from unittest.mock import patch


class Research:
    def __init__(self):
        self.calls = 0
        self.fail = False

    def ready(self):
        return None

    def discover(self, category=None, excluded=()):
        self.calls += 1
        if self.fail:
            raise SourceError('网络不可用')
        return [dict(platform=p, title=f'{p} 话题 {self.calls}-{n}', content='有来源的内容',
                     url=f'https://example.com/{p}/{self.calls}/{n}', event_key=f'{p}-{self.calls}-{n}',
                     published_at='2026-10-22T08:00:00+08:00', reason='来源依据', hook='开头建议')
                for p in ([category] if category else ['douyin', 'xiaohongshu', 'news']) for n in range(7)]


class BriefTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.now = datetime(2026, 10, 22, 10, tzinfo=ZoneInfo('Asia/Shanghai'))
        self.store = Store(self.tmp.name + '/test.db')
        self.research = Research()
        self.brief = Briefing(self.store, self.research, self.tmp.name, now=lambda: self.now)

    def tearDown(self):
        self.tmp.cleanup()

    def test_fixed_quotas_and_read_does_not_search(self):
        self.assertEqual(self.brief.report()['items'], [])
        result = self.brief.generate()
        self.assertEqual(Counter(i['platform'] for i in result['items']), {'douyin': 4, 'xiaohongshu': 3, 'news': 2})
        self.brief.generate()
        restored = Briefing(self.store, self.research, self.tmp.name, now=lambda: self.now)
        self.assertEqual(restored.report()['items'], result['items'])
        self.assertEqual(self.research.calls, 1)

    def test_replace_changes_only_one_slot_uses_backups_and_export(self):
        before = self.brief.generate()['items']
        after = self.brief.replace(before[0]['slot'], before[0]['id'])['items']
        self.assertNotEqual(after[0]['id'], before[0]['id'])
        self.assertEqual(after[1:], before[1:])
        self.assertEqual(self.research.calls, 1)
        book = load_workbook(io.BytesIO(self.brief.excel()))
        self.assertEqual(book['热点素材'].max_row, 10)
        self.assertEqual(book['热点素材']['A2'].value, after[0]['title'])
        self.assertTrue((Path(self.tmp.name) / 'reports' / 'real' / '2026-10-22.xlsx').is_file())
        with self.assertRaises(ValueError):
            self.brief.replace(before[0]['slot'], before[0]['id'])

    def test_exhausted_backups_researches_but_never_restores_rejected(self):
        row = self.brief.generate()['items'][0]
        seen = {row['id']}
        for _ in range(5):
            result = self.brief.replace(row['slot'], row['id'])
            row = result['items'][0]
            self.assertNotIn(row['id'], seen)
            seen.add(row['id'])
        self.assertEqual(self.research.calls, 2)

    def test_failed_replacement_preserves_all_items(self):
        self.brief.generate()
        for _ in range(3):
            row = self.brief.report()['items'][0]
            self.brief.replace(row['slot'], row['id'])
        before = self.brief.report()['items']
        self.research.fail = True
        result = self.brief.replace(before[0]['slot'], before[0]['id'])
        self.assertEqual(result['items'], before)
        self.assertIn('网络不可用', result['errors'][0]['message'])

    def test_dedup_missing_candidates_and_day_rollover(self):
        def limited(category=None, excluded=()):
            return [dict(platform=p, title='同一个事件', content='内容', url='https://example.com/shared',
                         event_key='shared', published_at='2026-10-22T08:00:00+08:00') for p in ['douyin', 'xiaohongshu', 'news']]
        self.research.discover = limited
        result = self.brief.generate()
        self.assertEqual(len(result['items']), 1)
        self.assertEqual(result['status'], 'partial')
        self.assertEqual(result['missing']['news'], 2)
        self.now += timedelta(days=1)
        self.assertEqual(self.brief.report()['items'], [])

    def test_api_replaces_exactly_one_and_export_reads_without_search(self):
        app = create_app(self.tmp.name, researcher=self.research, now=lambda: self.now)
        client = app.test_client()
        self.assertEqual(client.get('/api/brief').json['items'], [])
        self.assertEqual(client.get('/api/brief/export').status_code, 409)
        generated = client.post('/api/brief/generate', json={}).json
        self.assertEqual(len(generated['items']), 9)
        first = generated['items'][0]
        response = client.post('/api/brief/replace', json={'slot': first['slot'], 'expected_id': first['id']})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json['items'][1:], generated['items'][1:])
        self.assertEqual(client.post('/api/brief/replace', json={'slot': first['slot'], 'expected_id': first['id']}).status_code, 409)
        exported = client.get('/api/brief/export')
        self.assertEqual(load_workbook(io.BytesIO(exported.data))['热点素材'].max_row, 10)
        self.assertEqual(self.research.calls, 1)
        demo = client.get('/api/brief?demo=1').json
        self.assertTrue(demo['demo'])
        self.assertEqual(len(demo['items']), 9)
        self.assertFalse(client.get('/api/brief').json['demo'])

    def test_scheduler_only_attempts_daily_brief_once(self):
        self.research.fail = True
        tick(SimpleNamespace(now=self.brief.now), self.brief)
        tick(SimpleNamespace(now=self.brief.now), self.brief)
        self.assertEqual(self.research.calls, 1)
        self.research.fail = False
        self.brief.generate()
        self.assertEqual(self.research.calls, 2)

    def test_locked_excel_does_not_lose_saved_edition(self):
        with patch('hotdesk.briefs.os.replace', side_effect=PermissionError('locked')):
            result = self.brief.generate()
        self.assertEqual(len(result['items']), 9)
        self.assertIn('Excel', result['errors'][0]['message'])
        self.assertTrue(self.brief.excel())

    def test_rejected_url_cannot_return_with_a_different_model_event_key(self):
        initial = self.brief.generate()['items'][0]
        for _ in range(3):
            row = self.brief.report()['items'][0]
            self.brief.replace(row['slot'], row['id'])
        def recycled(category=None, excluded=()):
            return [dict(initial, event_key='changed-identifier', title='改名的同一个事件'),
                    dict(initial, event_key='genuinely-new-event', title='新的事件', url='https://example.com/new')]
        self.research.discover = recycled
        row = self.brief.report()['items'][0]
        result = self.brief.replace(row['slot'], row['id'])
        self.assertNotEqual(result['items'][0]['url'], initial['url'])


if __name__ == '__main__':
    unittest.main()
