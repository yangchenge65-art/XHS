import io
import json
import re
import sqlite3
import threading
from datetime import datetime, timedelta
from decimal import Decimal, ROUND_CEILING
from pathlib import Path
from urllib.parse import urlsplit
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

TZ = ZoneInfo('Asia/Shanghai')
PLATFORMS = ('douyin', 'xiaohongshu')
LABELS = {'douyin': '抖音', 'xiaohongshu': '小红书', 'news': '大众热门事件', 'openai': 'GPT-5.1'}


class SourceError(Exception):
    pass


def clock():
    return datetime.now(TZ)


def parse_time(value):
    try:
        if isinstance(value, (int, float)) or (isinstance(value, str) and value.isdigit()):
            stamp = float(value)
            return datetime.fromtimestamp(stamp / 1000 if stamp > 1e11 else stamp, TZ)
        dt = datetime.fromisoformat(str(value).replace('Z', '+00:00'))
        return dt.replace(tzinfo=TZ) if dt.tzinfo is None else dt.astimezone(TZ)
    except (ValueError, TypeError, OverflowError, OSError):
        return None


def safe_url(value):
    value = str(value or '')
    try:
        u = urlsplit(value)
        return value if u.scheme in ('http', 'https') and u.hostname and not u.username else ''
    except ValueError:
        return ''


class Store:
    def __init__(self, path):
        self.path = str(path)
        Path(path).parent.mkdir(parents=True, exist_ok=True)
        with self.connect() as db:
            db.executescript('''
                CREATE TABLE IF NOT EXISTS snapshots(
                    day TEXT, platform TEXT, identity TEXT, payload TEXT,
                    PRIMARY KEY(day, platform, identity));
                CREATE TABLE IF NOT EXISTS runs(
                    day TEXT, platform TEXT, updated TEXT, status TEXT, message TEXT,
                    PRIMARY KEY(day,platform));
                CREATE TABLE IF NOT EXISTS cache(key TEXT PRIMARY KEY, updated TEXT, payload TEXT);
                CREATE TABLE IF NOT EXISTS budget(month TEXT PRIMARY KEY, units INTEGER NOT NULL);
            ''')

    def connect(self):
        db = sqlite3.connect(self.path, timeout=30)
        db.row_factory = sqlite3.Row
        return db

    def reserve(self, now, cost):
        value = Decimal(str(cost))
        if not value.is_finite() or value < 0:
            raise ValueError('无效的请求预留费用')
        units = int((value * 10000).to_integral_value(rounding=ROUND_CEILING))
        month = now.astimezone(TZ).strftime('%Y-%m')
        with self.connect() as db:
            db.execute('BEGIN IMMEDIATE')
            db.execute('INSERT OR IGNORE INTO budget VALUES (?,0)', (month,))
            used = db.execute('SELECT units FROM budget WHERE month=?', (month,)).fetchone()[0]
            if used + units > 300000:
                return False
            db.execute('UPDATE budget SET units=units+? WHERE month=?', (units, month))
        return True

    def spent(self, now):
        with self.connect() as db:
            row = db.execute('SELECT units FROM budget WHERE month=?', (now.astimezone(TZ).strftime('%Y-%m'),)).fetchone()
        return row[0] / 10000 if row else 0

    def record(self, now, platform, status, message='', items=None):
        day = now.astimezone(TZ).date().isoformat()
        with self.connect() as db:
            if items is not None:
                for item in items:
                    identity = str(item.get('id') or item['title']).strip().casefold()
                    db.execute('INSERT OR REPLACE INTO snapshots VALUES(?,?,?,?)',
                               (day, platform, identity, json.dumps(item, ensure_ascii=False)))
            db.execute('INSERT OR REPLACE INTO runs VALUES(?,?,?,?,?)',
                       (day, platform, now.isoformat(), status, message))

    def runs(self, now):
        with self.connect() as db:
            rows = db.execute('SELECT * FROM runs WHERE day=?', (now.astimezone(TZ).date().isoformat(),)).fetchall()
        return {r['platform']: dict(r) for r in rows}

    def snapshots(self, now):
        with self.connect() as db:
            return db.execute('SELECT * FROM snapshots WHERE day BETWEEN ? AND ? ORDER BY day DESC',
                              ((now.astimezone(TZ).date() - timedelta(days=6)).isoformat(), now.astimezone(TZ).date().isoformat())).fetchall()

    def cache_get(self, key, now):
        with self.connect() as db:
            row = db.execute('SELECT * FROM cache WHERE key=?', (key,)).fetchone()
        if row and timedelta(0) <= now - parse_time(row['updated']) < timedelta(hours=6):
            return json.loads(row['payload'])

    def cache_set(self, key, now, data):
        with self.connect() as db:
            db.execute('INSERT OR REPLACE INTO cache VALUES(?,?,?)', (key, now.isoformat(), json.dumps(data, ensure_ascii=False)))


def decorate(item, platform, days=1, consecutive=1):
    item = dict(item)
    rank = item.get('rank')
    try:
        rank = int(rank) if rank else None
    except (ValueError, TypeError):
        rank = None
    score = min(100, days * 12 + max(0, 31 - (rank or 31)))
    reason = f'近七天记录 {days} 天；连续记录 {consecutive} 天'
    if rank:
        reason += f'；最近榜单第 {rank} 位'
    else:
        reason += '；相关内容，未证实上过热搜'
    title = str(item.get('title', ''))[:500]
    item.update(platform=platform, platform_name=LABELS[platform], title=title,
                content=str(item.get('content', title))[:10000], url=safe_url(item.get('url')),
                rank=rank, days=days, consecutive=consecutive, score=score, reason=reason,
                hook=f'“{title[:45]}”为什么值得关注？先看这一个细节。',
                hook_note='规则模板建议，使用前核实内容；不保证留存率。')
    return item


class Service:
    def __init__(self, store, provider, now=clock):
        self.store, self.provider, self.now = store, provider, now
        self.lock = threading.Lock()

    def request(self, platform, keyword=None):
        error = self.provider.ready(platform, 'search' if keyword else 'hot')
        if error:
            raise SourceError(error)
        if not self.store.reserve(self.now(), self.provider.cost):
            raise SourceError('本月 30 元本地预留额度不足，已停止新请求')
        return self.provider.fetch(platform, keyword)

    def collect(self):
        with self.lock:
            now = self.now()
            previous = self.store.runs(now)
            for platform in PLATFORMS:
                if previous.get(platform, {}).get('status') == 'ok':
                    continue
                try:
                    items = self.request(platform)
                    self.store.record(now, platform, 'ok', f'采集 {len(items)} 条', items)
                except SourceError as exc:
                    self.store.record(now, platform, 'error', str(exc))
            return self.report()

    def report(self):
        now = self.now()
        groups = {}
        for row in self.store.snapshots(now):
            item = json.loads(row['payload'])
            # Same topic with different supplier IDs is still counted once per day.
            key = (row['platform'], re.sub(r'[\s#＃]+', '', item['title']).casefold())
            if key not in groups:
                groups[key] = [item, set()]
            groups[key][1].add(row['day'])
        items = []
        for (platform, _), (item, dates) in groups.items():
            ordered = sorted(dates, reverse=True)
            consecutive = 1
            for a, b in zip(ordered, ordered[1:]):
                if (datetime.fromisoformat(a) - datetime.fromisoformat(b)).days != 1:
                    break
                consecutive += 1
            enriched = decorate(item, platform, len(dates), consecutive)
            enriched.update(observed_at=ordered[0], kind='榜单记录')
            items.append(enriched)
        items.sort(key=lambda i: (-i['score'], i['title']))
        runs = self.store.runs(now)
        errors = [dict(platform=p, message=runs[p]['message']) for p in PLATFORMS if runs.get(p, {}).get('status') == 'error']
        return dict(items=items, errors=errors, runs=runs, status='partial' if errors and items else 'error' if errors else 'ok',
                    generated_at=now.isoformat(), cached=False, demo=False,
                    coverage='近七个自然日的本地榜单记录；优先显示持续出现和排名靠前的话题。出现次数不是讨论量。')

    def search(self, keyword):
        keyword = str(keyword).strip()
        if not 1 <= len(keyword) <= 80:
            raise ValueError('请输入 1–80 个字的关键词')
        now = self.now()
        key = keyword.casefold()
        with self.lock:
            cached = self.store.cache_get(key, now)
            if cached is not None:
                cached['cached'] = True
                return self.filter_window(cached, now)
            result = dict(items=[], errors=[], keyword=keyword, cached=False, demo=False, generated_at=now.isoformat(),
                          coverage='本地近七天榜单记录 + 近 168 小时相关内容。每个平台仅搜索第一页；不代表完整历史热搜。相关内容缺少可验证发布时间时不纳入。')
            result['items'].extend(i for i in self.report()['items']
                                   if key in (i['title'] + ' ' + i['content']).casefold())
            for platform in PLATFORMS:
                try:
                    for item in self.request(platform, keyword):
                        dt = parse_time(item.get('published_at'))
                        if not dt or not now - timedelta(days=7) <= dt <= now:
                            continue
                        enriched = decorate(item, platform)
                        enriched.update(kind='相关内容', published_at=dt.isoformat(), observed_at=now.isoformat())
                        # Search rank/likes are not a historical hot-list position.
                        enriched['rank'] = None
                        enriched['reason'] = '近七天关键词相关内容；未证实上过热搜'
                        result['items'].append(enriched)
                except SourceError as exc:
                    result['errors'].append(dict(platform=platform, message=str(exc)))
            result['status'] = 'partial' if result['errors'] and result['items'] else 'error' if result['errors'] else 'ok'
            if not result['errors']:
                self.store.cache_set(key, now, result)
            return result

    @staticmethod
    def filter_window(result, now):
        result['items'] = [i for i in result['items']
                           if (dt := parse_time(i.get('observed_at') if i.get('kind') == '榜单记录' else i.get('published_at')))
                           and now - timedelta(days=7) <= dt <= now]
        return result


def export_excel(report, demo=False):
    book = Workbook()
    sheet = book.active
    sheet.title = '热点素材'
    sheet.append(['热点关键词', '热点具体内容', '热点相关链接', '平台', '内容类型', '发布时间 / 最近记录',
                  '近七天记录天数', '筛选依据', '开头钩子建议', '数据模式'])
    for item in report['items']:
        values = [item['title'], item['content'], item['url'], item['platform_name'], item['kind'],
                  item.get('published_at') or item.get('observed_at', ''), item['days'], item['reason'],
                  item['hook'], '示例数据' if demo else '真实数据源']
        sheet.append(values)
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                # Explicit string cells avoid Excel formula injection without changing text.
                cell.data_type = 's'
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = PatternFill('solid', fgColor='166E53')
        cell.font = Font(color='FFFFFF', bold=True)
    for col, width in zip('ABCDEFGHIJ', [30, 65, 48, 12, 15, 28, 18, 55, 60, 15]):
        sheet.column_dimensions[col].width = width
    notes = book.create_sheet('采集说明')
    for row in [('数据模式', '示例数据：仅用于体验，不是真实热点' if demo else '真实数据源；仅对已获取的数据负责'),
                ('生成时间', report.get('generated_at', '')), ('覆盖范围', report.get('coverage', '')),
                ('整理方式', report.get('engine', '来源内容 / 本地规则')),
                ('查询状态', report.get('status', '')), ('提醒', '钩子为创作建议，需人工核实，不保证留存率。')]:
        notes.append(row)
    for error in report.get('errors', []):
        notes.append((LABELS.get(error['platform'], error['platform']), error['message']))
    for row in notes:
        for cell in row:
            cell.data_type = 's'
    notes.column_dimensions['A'].width = 18
    notes.column_dimensions['B'].width = 100
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()
